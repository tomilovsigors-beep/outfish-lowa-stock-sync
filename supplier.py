import io, re
from collections import defaultdict
from openpyxl import load_workbook

SIZE_RE = re.compile(r"^\s*(\d+(?:[.,]\d+)?)")

def norm_text(v):
    return re.sub(r"\s+", " ", str(v or "").strip()).lower()

def norm_model(v):
    s = str(v or "").replace("®", "").strip()
    s = re.sub(r"\s+", " ", s)
    return s

def norm_size(v):
    if v is None: return None
    s = str(v).strip().replace(",", ".").replace("-", ".")
    m = SIZE_RE.match(s)
    if not m: return None
    n = float(m.group(1))
    return str(int(n)) if n.is_integer() else ("%.1f" % n)

def qty(v):
    if v is None or v == "": return 0
    if isinstance(v, (int,float)): return int(v)
    m = re.match(r"\s*(-?\d+)", str(v))
    return int(m.group(1)) if m else 0

def parse_xlsx_bytes(data: bytes):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb["LOWA"]
    rows = list(ws.iter_rows(values_only=True))
    out = defaultdict(lambda: defaultdict(int))
    meta = {}
    i = 0
    while i < len(rows):
        row = rows[i]
        model_cell = row[1] if len(row) > 1 else None
        # Header rows have size labels from column C onward.
        if model_cell and any(norm_size(x) for x in row[2:25] if isinstance(x, str) and "(" in x):
            model = norm_model(model_cell)
            sizes = [norm_size(x) for x in row[2:25]]
            i += 1
            while i < len(rows):
                r = rows[i]
                color_cell = r[1] if len(r) > 1 else None
                # stop at blank separator or next header
                if color_cell is None and not any(x not in (None, "") for x in r[:25]):
                    break
                if color_cell and any(norm_size(x) for x in r[2:25] if isinstance(x, str) and "(" in x):
                    break
                color_raw = str(color_cell or "").strip()
                if not color_raw:
                    i += 1; continue
                lithuania = "*" in color_raw
                color_clean = color_raw.replace("*", "").strip()
                wide = bool(re.search(r"\bW\s*$", color_clean, re.I))
                if wide:
                    color_clean = re.sub(r"\bW\s*$", "", color_clean, flags=re.I).strip()
                color = norm_text(color_clean)
                width = "wide" if wide else "standard"
                key = (model, color, width)
                for idx, size in enumerate(sizes, start=2):
                    if size is None or idx >= len(r):
                        continue
                    out[key][size] += qty(r[idx])
                meta[key] = {"model": model, "color": color, "width": width}
                i += 1
            continue
        i += 1
    return [{**meta[k], "stock": dict(sorted(v.items(), key=lambda kv: float(kv[0])))} for k,v in out.items()]
